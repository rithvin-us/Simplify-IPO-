"""POST /parse — extract raw text from an uploaded document.

Uses pymupdf / python-docx / openpyxl when installed; otherwise decodes the
bytes as text. Never hard-fails on a missing optional parser.
"""
from __future__ import annotations

import io

from fastapi import APIRouter, File, Form, UploadFile

router = APIRouter(tags=["parse"])


def _parse_pdf(data: bytes) -> str | None:
    try:
        import fitz  # pymupdf
    except ImportError:
        return None
    doc = fitz.open(stream=data, filetype="pdf")
    return "\n".join(page.get_text() for page in doc)


def _parse_docx(data: bytes) -> str | None:
    try:
        import docx  # python-docx
    except ImportError:
        return None
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _parse_xlsx(data: bytes) -> str | None:
    try:
        import openpyxl
    except ImportError:
        return None
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return (text, method). method notes which parser produced the text."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _parse_pdf(data)
        if text is not None:
            return text, "pymupdf"
    elif name.endswith(".docx"):
        text = _parse_docx(data)
        if text is not None:
            return text, "python-docx"
    elif name.endswith((".xlsx", ".xlsm")):
        text = _parse_xlsx(data)
        if text is not None:
            return text, "openpyxl"
    # Fallback: treat as UTF-8 text (txt, csv, md, or parser lib absent).
    return data.decode("utf-8", errors="replace"), "raw-text"


@router.post("/parse")
async def parse(file: UploadFile = File(...), filename: str = Form(default="")) -> dict:
    data = await file.read()
    name = filename or file.filename or "upload"
    text, method = extract_text(name, data)
    return {
        "filename": name,
        "method": method,
        "chars": len(text),
        "text": text,
    }
